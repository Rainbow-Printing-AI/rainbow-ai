"""
pb_render.py - renderers for the Rainbow price book (used by pricebook.py).

  book_pdf        internal master price book (letter, binder-ready, DRAFT watermark while cards are open)
  conflicts_pdf   printable conflict / question sheet for mark-up and scan-back
  xlsx            filterable workbook for the owners
  html            single-file offline search page (also the source of the published web page)
  mirror_md       compact text copy for the Claude project

Only needs reportlab, openpyxl and Pillow (for reportlab's image handling).
"""
import csv, datetime as dt, html as H, json, os, re
from collections import OrderedDict, defaultdict

import pricebook as pb

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Table, TableStyle,
                                KeepTogether, Image, Flowable, PageBreak, CondPageBreak)
from reportlab.pdfgen import canvas as rl_canvas

INK = colors.HexColor("#1d2024")
GREY = colors.HexColor("#6b7178")
LIGHT = colors.HexColor("#e9ecef")
RULE = colors.HexColor("#b9bfc6")
CYAN = colors.HexColor("#0088c7")
MAG = colors.HexColor("#c2185b")
AMBER = colors.HexColor("#b26a00")

def S(name, **kw):
    base = dict(fontName="Helvetica", fontSize=9, leading=11.5, textColor=INK)
    base.update(kw)
    return ParagraphStyle(name, **base)

ST = {
    "h1": S("h1", fontName="Helvetica-Bold", fontSize=22, leading=26),
    "h2": S("h2", fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=4, spaceAfter=4),
    "h3": S("h3", fontName="Helvetica-Bold", fontSize=10, leading=12.5),
    "body": S("body"),
    "small": S("small", fontSize=7.5, leading=9.2, textColor=GREY),
    "cell": S("cell", fontSize=8.3, leading=10),
    "cellb": S("cellb", fontName="Helvetica-Bold", fontSize=8.3, leading=10),
    "cellr": S("cellr", fontSize=8.3, leading=10, alignment=TA_RIGHT),
    "price": S("price", fontName="Helvetica-Bold", fontSize=8.6, leading=10, alignment=TA_RIGHT),
    "note": S("note", fontSize=7, leading=8.4, textColor=GREY),
    "flag": S("flag", fontName="Helvetica-Bold", fontSize=6.8, leading=8, textColor=AMBER),
    "cardid": S("cardid", fontName="Helvetica-Bold", fontSize=17, leading=19),
    "cardtitle": S("cardtitle", fontName="Helvetica-Bold", fontSize=10.5, leading=12.5),
    "ctx": S("ctx", fontSize=8.4, leading=10.4),
    "cap": S("cap", fontSize=6.8, leading=8, textColor=GREY),
    "line": S("line", fontSize=9, leading=11),
    "prop": S("prop", fontName="Helvetica-Bold", fontSize=9.5, leading=11, alignment=TA_RIGHT),
}

def esc(s):
    return H.escape(str(s or ""), quote=False)

def P(text, style="cell"):
    return Paragraph(text, ST[style] if isinstance(style, str) else style)

class Box(Flowable):
    """an empty checkbox"""
    def __init__(self, size=0.17 * inch, filled=False):
        super().__init__()
        self.size, self.filled = size, filled
        self.width = self.height = size
    def draw(self):
        self.canv.setLineWidth(1)
        self.canv.setStrokeColor(INK)
        self.canv.rect(0, 0, self.size, self.size, stroke=1, fill=0)

class WriteIn(Flowable):
    """an empty write-in box, optionally with a $ prefix"""
    def __init__(self, width=1.2 * inch, height=0.3 * inch, prefix="$"):
        super().__init__()
        self.width, self.height, self.prefix = width, height, prefix
    def draw(self):
        c = self.canv
        c.setLineWidth(0.9); c.setStrokeColor(INK)
        c.rect(0, 0, self.width, self.height, stroke=1, fill=0)
        if self.prefix:
            c.setFont("Helvetica", 9); c.setFillColor(GREY)
            c.drawString(4, self.height / 2 - 3.2, self.prefix)

class NumberedCanvas(rl_canvas.Canvas):
    """defers page decoration so every page can say 'page N of M'"""
    def __init__(self, *a, **kw):
        self._decorate = kw.pop("decorate", None)
        super().__init__(*a, **kw)
        self._saved = []
    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()
    def save(self):
        n = len(self._saved)
        for st in self._saved:
            self.__dict__.update(st)
            if self._decorate:
                self._decorate(self, self._pageNumber, n)
            super().showPage()
        super().save()

def canvasmaker(decorate):
    def make(*a, **kw):
        kw["decorate"] = decorate
        return NumberedCanvas(*a, **kw)
    return make

def cmyk_marks(c, x, y, s=5):
    for i, col in enumerate(("#00a0dc", "#d4147f", "#f2c500", "#1d2024")):
        c.setFillColor(colors.HexColor(col))
        c.rect(x + i * (s + 1.5), y, s, s, stroke=0, fill=1)

# ============================================================== conflict sheet
def conflicts_pdf(root, include=None):
    doc_json = pb.load_cards()
    ver = doc_json.get("version", 1)
    cards = []
    for c in doc_json["cards"]:
        if c["status"] != "open" or (include and c["id"] not in include):
            continue
        c = dict(c)
        c["lines"] = [ln for ln in c.get("lines", []) if "answer" not in ln]
        if c.get("answer_option") is not None:
            c["options"] = []
        if not c["lines"] and not c["options"] and not c.get("list_items"):
            continue
        cards.append(c)
    crops = os.path.join(root, "source", "crops")
    out = os.path.join(root, "build", f"Rainbow_Price_Conflicts_v{ver}.pdf")
    printed = dt.date.today().isoformat()
    W, Hh = letter
    M = 0.5 * inch
    avail = W - 2 * M - 0.24 * inch

    def decorate(c, n, total):
        c.saveState()
        c.setFont("Helvetica-Bold", 8.5); c.setFillColor(INK)
        c.drawString(M, Hh - 0.36 * inch, f"RAINBOW PRICE BOOK  |  CONFLICT SHEET v{ver}")
        c.setFont("Helvetica", 8)
        c.drawRightString(W - M, Hh - 0.36 * inch, f"printed {printed}   |   page {n} of {total}   |   PBCS-v{ver}-p{n}")
        cmyk_marks(c, M + 2.72 * inch, Hh - 0.37 * inch, 5)
        c.setStrokeColor(INK); c.setLineWidth(0.6)
        c.line(M, Hh - 0.43 * inch, W - M, Hh - 0.43 * inch)
        # per-page blanket OK
        y = 0.34 * inch
        c.setLineWidth(1); c.rect(M, y - 2, 0.17 * inch, 0.17 * inch)
        c.setFont("Helvetica-Bold", 8.5)
        c.drawString(M + 0.26 * inch, y + 1.5,
                     "OK every proposal on this page that I did not mark otherwise.")
        c.setFont("Helvetica", 8); c.setFillColor(GREY)
        c.drawRightString(W - M, y + 1.5, "Initials ______   Date ________")
        c.restoreState()

    story = []
    story.append(P("Price conflicts &amp; questions", "h1"))
    story.append(Spacer(1, 4))
    n_lines = sum(len(c.get("lines", [])) or 1 for c in cards)
    intro = (f"<b>{len(cards)} cards</b> from the 2026-09-30 scans. Each card shows what the sheets say, a crop of the "
             "actual scan, and my proposed answer.")
    story.append(P(intro, "body"))
    story.append(Spacer(1, 5))
    how = [
        "<b>Agree?</b> Tick <b>OK</b> on the line, or tick the box at the bottom of the page to OK every proposal you didn't mark.",
        "<b>Disagree?</b> Write the right price in the <b>$ box</b>. On option cards, tick the option you want.",
        "<b>Drop an item?</b> Tick <b>Remove</b> where offered, or write 'remove' in the $ box.",
        "<b>Not sure yet?</b> Leave it blank. Blank = still open; it rolls to the next sheet.",
        "Write in dark pen, one digit per character, with the decimal point. Then scan all pages into the Rainbow folder.",
    ]
    story.append(Table([[P("&#8226;", "body"), P(t, "body")] for t in how], colWidths=[0.15 * inch, avail],
                       style=[("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5), ("TOPPADDING", (0, 0), (-1, -1), 0)]))
    story.append(Spacer(1, 6))
    rules = doc_json.get("rules", {})
    rt = [[P(f"<b>{k}</b>", "small"), P(esc(v), "small")] for k, v in rules.items()]
    story.append(Table([[P("<b>How I proposed prices</b>", "small"), ""]] + rt, colWidths=[0.4 * inch, avail - 0.25 * inch],
                       style=[("SPAN", (0, 0), (1, 0)), ("BOX", (0, 0), (-1, -1), 0.5, RULE),
                              ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f5f6f7")),
                              ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.2),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2)]))
    story.append(Spacer(1, 10))

    def crop_flowables(card):
        out = []
        items = []
        for cr in card.get("crops", []):
            p = os.path.join(crops, cr["file"])
            if not os.path.exists(p):
                continue
            from PIL import Image as PILImage
            with PILImage.open(p) as im:
                iw, ih = im.size
            items.append((p, iw / ih, cr["caption"], cr.get("h")))
        wide = [x for x in items if x[1] > 3.2]
        norm = [x for x in items if x[1] <= 3.2]
        for p, ar, cap, hh in wide:
            h = hh * inch if hh else (0.42 * inch if ar > 8 else 0.6 * inch if ar > 5 else 0.9 * inch)
            w = h * ar
            if w > avail:
                w, h = avail, avail / ar
            out.append(Table([[Image(p, width=w, height=h)], [P(esc(cap), "cap")]], colWidths=[avail],
                             style=[("LEFTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 1),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
        for i in range(0, len(norm), 2):
            row = norm[i:i + 2]
            colw = avail / len(row) - 4
            cells, caps = [], []
            for p, ar, cap, hh in row:
                maxh = hh * inch if hh else (1.5 * inch if len(row) == 2 else 1.75 * inch)
                w = min(colw, maxh * ar); h = w / ar
                cells.append(Image(p, width=w, height=h)); caps.append(P(esc(cap), "cap"))
            out.append(Table([cells, caps], colWidths=[avail / len(row)] * len(row),
                             style=[("LEFTPADDING", (0, 0), (-1, -1), 0), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                    ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
        return out

    kind_label = {"conflict": "CONFLICT", "verify": "CHECK MY READING", "question": "QUESTION", "gap": "MISSING PRICE"}
    for card in cards:
        inner = []
        head = Table([[P(card["id"], "cardid"),
                       [P(f"<font size=7 color='#6b7178'>{esc(card['section']).upper()}  |  {kind_label.get(card['kind'], '')}</font>", "small"),
                        P(esc(card["title"]), "cardtitle")]]],
                     colWidths=[0.62 * inch, avail - 0.62 * inch],
                     style=[("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 2)])
        inner.append(head)
        for t in card.get("context", []):
            inner.append(P(esc(t), "ctx"))
        inner.append(Spacer(1, 3))
        inner += crop_flowables(card)
        if card.get("list_items"):
            li = card["list_items"]
            ncol = 3
            per = -(-len(li) // ncol)
            cols = [li[i * per:(i + 1) * per] for i in range(ncol)]
            data = []
            for i in range(per):
                data.append([P(esc(cols[j][i]) if i < len(cols[j]) else "", "note") for j in range(ncol)])
            inner.append(Table(data, colWidths=[avail / ncol] * ncol,
                               style=[("LEFTPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 0.6),
                                      ("BOTTOMPADDING", (0, 0), (-1, -1), 0.6)]))
            inner.append(Spacer(1, 3))
        if card.get("lines"):
            anyrem = any(l.get("allow_remove") for l in card["lines"])
            data = [[P("<font size=7 color='#6b7178'>ITEM</font>", "small"), P("<font size=7 color='#6b7178'>PROPOSED</font>", "small"),
                     P("<font size=7 color='#6b7178'>OK</font>", "small"), P("<font size=7 color='#6b7178'>OR USE</font>", "small")]
                    + ([P("<font size=7 color='#6b7178'>REMOVE</font>", "small")] if anyrem else [])]
            for ln in card["lines"]:
                prop = ln.get("proposed")
                if prop in (None, ""):
                    ptxt, okcell = "<font color='#b26a00'>no proposal</font>", ""
                elif prop == "retire":
                    ptxt, okcell = "discontinue", Box()
                else:
                    try:
                        ptxt = f"${float(prop):,.2f}"
                    except ValueError:
                        ptxt = esc(prop)
                    okcell = Box()
                row = [P(f"<b>{ln['key']}.</b> {esc(ln['label'])}", "line"), P(ptxt, "prop"), okcell,
                       WriteIn(1.25 * inch, 0.3 * inch, "Keep $" if prop == "retire" else "$")]
                if anyrem:
                    row.append(Box() if ln.get("allow_remove") else "")
                data.append(row)
            cw = [avail - 3.1 * inch - (0.55 * inch if anyrem else 0), 0.95 * inch, 0.45 * inch, 1.7 * inch] + ([0.55 * inch] if anyrem else [])
            t = Table(data, colWidths=cw, rowHeights=[0.2 * inch] + [0.4 * inch] * len(card["lines"]))
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 1), (-1, -1), 0.4, LIGHT),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 3), ("ALIGN", (2, 0), (2, -1), "CENTER")]
                                  + ([("ALIGN", (4, 0), (4, -1), "CENTER")] if anyrem else [])))
            inner.append(t)
        if card.get("options"):
            po = card.get("proposed_option")
            data = []
            for i, o in enumerate(card["options"]):
                lab = f"<b>{i + 1}.</b> {esc(o)}"
                if po is not None and i == po:
                    lab += "  <font color='#0088c7'><b>(proposed)</b></font>"
                data.append([Box(), P(lab, "line")])
            inner.append(Table(data, colWidths=[0.3 * inch, avail - 0.3 * inch],
                               style=[("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3.2),
                                      ("BOTTOMPADDING", (0, 0), (-1, -1), 3.2), ("LEFTPADDING", (0, 0), (-1, -1), 2)]))
        if card.get("extra_question"):
            inner.append(Spacer(1, 3))
            inner.append(P(esc(card["extra_question"]).replace("[ ]", "[&nbsp;&nbsp;&nbsp;]"), "line"))
        inner.append(Spacer(1, 5))
        inner.append(P("Notes: ________________________________________________________________________________________________", "line"))
        box = Table([[inner]], colWidths=[W - 2 * M],
                    style=[("BOX", (0, 0), (-1, -1), 1.1, INK), ("LEFTPADDING", (0, 0), (-1, -1), 8),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 6),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 7)])
        story.append(KeepTogether([box, Spacer(1, 9)]))

    doc = BaseDocTemplate(out, pagesize=letter, leftMargin=M, rightMargin=M, topMargin=0.58 * inch,
                          bottomMargin=0.62 * inch, title=f"Rainbow Price Conflicts v{ver}", author="Rainbow price book")
    frame = Frame(M, 0.62 * inch, W - 2 * M, Hh - 1.2 * inch, id="f", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([PageTemplate(id="p", frames=[frame])])
    doc.build(story, canvasmaker=canvasmaker(decorate))
    return out

# ============================================================== master book
def live_rows(rows):
    return [r for r in rows if r.get("status") in pb.LIVE]

def sort_key(r):
    d = pb.size_dims(r.get("size", ""))
    area = d[0] * d[1] if d else 10 ** 6
    fam = r.get("family") or r["item"].lower()
    q = int(r["qty_min"]) if r.get("qty_min", "").isdigit() else 0
    if r.get("category") in ("Bindery & Finishing", "Labor & Design", "Computer & File Services", "Scanning", "Special Inks"):
        return (r["item"].lower(), fam, q, r["id"])
    if r.get("category") in ("NCR Forms", "Envelopes", "Business Cards", "Labels"):
        return (fam, area, "addl" in r["id"], q, r["id"])
    walk = 0 if "walk-in" in r.get("variant", "") else 1
    return (area, r.get("sides", ""), walk, fam, r.get("unit") == "box", q, r["id"])

CAT_NOTES = {
    "Copies - B/W": "Charged per side (total page count). Double-sided = 2 sides; no duplex discount, no quantity breaks.",
    "Copies - Color": "Charged per side (total page count). Double-sided = 2 sides; no duplex discount. Lower rate from 21 sides up.",
    "Large Format - B/W": "No separate account pricing - account customers pay these prices.",
    "Labels": "Avery 5160 white, 30 per sheet. Includes type/set-up. Turnaround same or next day. Prices do not include sales tax.",
    "Business Cards": "Standard US size 3.5x2. Other sizes: special quote. Prices do not include artwork/layout. "
                      "Best file: high-res PDF with crop marks and bleeds. In-house bleed jobs run 8-up (20% fewer) + $5 cutting.",
    "Signs & Banners (4over)": "Ordered through 4over (Mark). Approx. 1 week. Artwork/layout extra.",
    "Photo Prints": "Papers: glossy, semi-gloss/matte, heavyweight matte.",
    "NCR Forms": "Trade (outside printer) pricing; retail = listed x 2.5 (handwritten on the sheet) - see card C11.",
    "Envelopes": "Trade (outside printer) pricing; retail = listed x 2.5 - see card C11.",
    "Pens & Markers": "Current = the red-pen sheet. 'box' = second price column (box/dozen), never updated - see card C22.",
    "Computer & File Services": "File printing charges are plus the copy charge.",
}

def book_pdf(root, draft="auto"):
    rows = pb.load_prices()
    contacts = pb.load_contacts()
    cards = pb.load_cards()
    open_cards = [c for c in cards["cards"] if c["status"] != "resolved"]
    is_draft = (draft == "yes") or (draft == "auto" and bool(open_cards))
    stamp = dt.date.today().isoformat()
    out = os.path.join(root, "build", f"Rainbow_PriceBook_{stamp}.pdf")
    W, Hh = letter
    M = 0.55 * inch
    avail = W - 2 * M
    live = live_rows(rows)
    bycat = OrderedDict((c, []) for c in pb.CATEGORY_ORDER)
    for r in live:
        bycat.setdefault(r["category"], []).append(r)

    def decorate(c, n, total):
        c.saveState()
        if is_draft:
            c.setFont("Helvetica-Bold", 90); c.setFillColor(colors.Color(0.8, 0.1, 0.1, alpha=0.07))
            c.translate(W / 2, Hh / 2); c.rotate(40); c.drawCentredString(0, -30, "DRAFT"); c.rotate(-40)
            c.translate(-W / 2, -Hh / 2)
        c.setFillColor(GREY); c.setFont("Helvetica", 7.3)
        left = f"Rainbow Price Book  |  generated {stamp} from data/prices.csv"
        if is_draft:
            left += f"  |  DRAFT: {len(open_cards)} open questions"
        c.drawString(M, 0.38 * inch, left)
        c.drawRightString(W - M, 0.38 * inch, f"page {n} of {total}")
        cmyk_marks(c, W - M - 0.9 * inch, 0.37 * inch, 4)
        c.restoreState()

    def build(section_pages):
        story = []
        story.append(Spacer(1, 0.3 * inch))
        story.append(P("Rainbow Printing &amp; Office Supply", "small"))
        story.append(P("Price Book", "h1"))
        story.append(P(f"Internal edition  |  generated {stamp}", "body"))
        story.append(Spacer(1, 8))
        if is_draft:
            story.append(Table([[P(f"<b>DRAFT.</b> {len(open_cards)} questions are open on the conflict sheet. "
                                   "Prices marked <font color='#b26a00'><b>PENDING Cxx</b></font> are my proposal from the scans; "
                                   "<font color='#b26a00'><b>NO PRICE</b></font> means the sheets don't give a usable one.", "body")]],
                               colWidths=[avail], style=[("BOX", (0, 0), (-1, -1), 0.8, AMBER),
                                                         ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#fff6e8")),
                                                         ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                                         ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.append(Spacer(1, 8))
        story.append(P("Prices exclude sales tax unless stated. Sizes are in inches, short side first (24x36 = 36x24). "
                       "The master data is Rainbow/PriceBook/data/prices.csv; ask Claude to look up or change a price, "
                       "and this book is regenerated.", "body"))
        story.append(Spacer(1, 10))
        story.append(P("Contents", "h2"))
        toc = []
        for cat, rs in bycat.items():
            if rs:
                toc.append([P(esc(cat), "cell"), P(str(len(rs)) + " prices", "note"), P(str(section_pages.get(cat, "")), "cellr")])
        toc.append([P("Phone list", "cell"), P(str(sum(1 for c in contacts if c.get("status") != "retired")) + " entries", "note"), P(str(section_pages.get("_contacts", "")), "cellr")])
        toc.append([P("Keyword index", "cell"), P("", "note"), P(str(section_pages.get("_index", "")), "cellr")])
        story.append(Table(toc, colWidths=[3.2 * inch, 1.5 * inch, 0.6 * inch],
                           style=[("LINEBELOW", (0, 0), (-1, -1), 0.3, LIGHT), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]))
        for cat, rs in bycat.items():
            if not rs:
                continue
            story.append(PageBreak() if cat in ("Pens & Markers",) else CondPageBreak(1.6 * inch))
            mk = SectionMark(cat)
            story.append(mk)
            story.append(section_header(cat, avail))
            if CAT_NOTES.get(cat):
                story.append(P(esc(CAT_NOTES[cat]), "note"))
                story.append(Spacer(1, 3))
            rs = sorted(rs, key=sort_key)
            if cat == "Pens & Markers":
                story += pens_table(rs, avail)
            else:
                story.append(price_table(rs, avail))
            story.append(Spacer(1, 8))
        story.append(PageBreak())
        story.append(SectionMark("_contacts"))
        story.append(section_header("Phone list", avail))
        story.append(P("Logins, passwords and tax IDs are kept only in private/logins.csv - not printed here.", "note"))
        story.append(Spacer(1, 3))
        story.append(contacts_table(contacts, avail))
        story.append(PageBreak())
        story.append(SectionMark("_index"))
        story.append(section_header("Keyword index", avail))
        story.append(index_table(live, section_pages, avail))
        return story

    class SectionMark(Flowable):
        def __init__(self, key):
            super().__init__(); self.key = key; self.width = self.height = 0
        def draw(self):
            pages_seen.setdefault(self.key, self.canv.getPageNumber())

    pages_seen = {}
    frame = Frame(M, 0.6 * inch, avail, Hh - 1.1 * inch, id="f", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    for _pass in range(2):
        prev = dict(pages_seen)
        pages_seen.clear()
        doc = BaseDocTemplate(out, pagesize=letter, leftMargin=M, rightMargin=M, topMargin=0.5 * inch, bottomMargin=0.6 * inch,
                              title="Rainbow Price Book", author="Rainbow Printing & Office Supply")
        doc.addPageTemplates([PageTemplate(id="p", frames=[frame])])
        doc.build(build(prev), canvasmaker=canvasmaker(decorate))
    return out

def section_header(title, avail):
    t = Table([[P(f"<font color='white'><b>{esc(title)}</b></font>", S("sh", fontName="Helvetica-Bold", fontSize=11.5, leading=14, textColor=colors.white))]],
              colWidths=[avail], style=[("BACKGROUND", (0, 0), (-1, -1), INK), ("LEFTPADDING", (0, 0), (-1, -1), 7),
                                        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)])
    return t

def row_flag(r):
    if r["status"] == "pending":
        return f"PENDING {r.get('cards', '')}"
    if r["status"] == "unpriced":
        return f"NO PRICE {r.get('cards', '')}"
    return ""

def show_note(r):
    n = r.get("notes", "")
    n = re.sub(r"Trade sheet lists \$[\d.,]+; price = x2\.5 per handwritten note\.\s*", "", n)
    if r.get("cost"):
        n = (f"trade ${float(r['cost']):,.2f}. " + n).strip()
    return n

def price_table(rs, avail):
    head = ["Item", "Details", "Size", "Qty", "Price", "Unit", ""]
    data = [[P(f"<font size=6.8 color='#6b7178'>{h.upper()}</font>", "small") for h in head]]
    style = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, 0), 0.6, INK),
             ("LINEBELOW", (0, 1), (-1, -1), 0.25, LIGHT), ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
             ("RIGHTPADDING", (0, 0), (-1, -1), 2.5), ("TOPPADDING", (0, 0), (-1, -1), 1.8), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8)]
    prev_item = None
    for i, r in enumerate(rs, start=1):
        details = esc(", ".join(x for x in (r.get("variant"), r.get("stock") if r.get("stock") not in ("", "20# bond") else "",
                                              r.get("color") if r.get("color") in ("Black ink", "Color ink") else "") if x))
        if r.get("min_charge"):
            details += f" <font color='#6b7178'>(min {pb.money(r['min_charge'])})</font>"
        n = show_note(r)
        if n:
            details += f"<br/><font size=6.6 color='#6b7178'>{esc(n)}</font>"
        item = esc(r["item"]) if r["item"] != prev_item else "<font color='#b9bfc6'>″</font>"
        prev_item = r["item"]
        price = pb.fmt_price(r) if r.get("price") else ("<font color='#b26a00'>no price</font>" if r["status"] != "retired" else "-")
        if not r.get("price") and "estimate" in r.get("notes", "").lower():
            price = "estimate"
        data.append([P(item, "cellb" if item and not item.startswith("<font") else "cell"), P(details, "cell"),
                     P(esc(r.get("size", "")), "cell"), P(esc(pb.qty_label(r)), "cell"), P(price, "price"),
                     P(esc(pb.unit_label(r).lstrip("/")), "cell"), P(esc(row_flag(r)), "flag")])
        if r["status"] in ("pending", "unpriced"):
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#fff8ec")))
    cw = [1.55 * inch, 2.5 * inch, 0.62 * inch, 0.72 * inch, 0.72 * inch, 0.5 * inch, 0.79 * inch]
    scale = avail / sum(cw)
    return Table(data, colWidths=[w * scale for w in cw], style=style, repeatRows=1)

def pens_table(rs, avail):
    each = [r for r in rs if r["unit"] == "each"]
    box = {r["family"]: r for r in rs if r["unit"] == "box"}
    each.sort(key=lambda r: r["item"].lower())
    rowsdata = []
    for r in each:
        b = box.get(r["family"])
        name = esc(r["item"]) + (f" <font size=6.5 color='#6b7178'>{esc(r['sku'])}</font>" if r.get("sku") else "")
        fl = r.get("cards", "") if r["status"] != "active" else ""
        rowsdata.append((name, pb.money(r["price"]) if r.get("price") else "-", pb.money(b["price"]) if b else "", fl, r["status"]))
    half = -(-len(rowsdata) // 2)
    cols = [rowsdata[:half], rowsdata[half:]]
    hdr = [P("<font size=6.8 color='#6b7178'>ITEM</font>", "small"), P("<font size=6.8 color='#6b7178'>EACH</font>", "small"),
           P("<font size=6.8 color='#6b7178'>BOX*</font>", "small"), P("", "small")]
    out_rows = [hdr + [""] + hdr]
    style = [("LINEBELOW", (0, 0), (-1, 0), 0.6, INK), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
             ("TOPPADDING", (0, 0), (-1, -1), 1.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2)]
    for i in range(half):
        line = []
        for j, col in enumerate(cols):
            if i < len(col):
                name, e, b, fl, st = col[i]
                line += [P(name, S("pn", fontSize=7.6, leading=9)), P(e, S("pe", fontName="Helvetica-Bold", fontSize=7.8, leading=9, alignment=TA_RIGHT)),
                         P(b, S("pb", fontSize=7.2, leading=9, alignment=TA_RIGHT, textColor=GREY)), P(esc(fl), "flag")]
                if st != "active":
                    c0 = 0 if j == 0 else 5
                    style.append(("BACKGROUND", (c0, i + 1), (c0 + 3, i + 1), colors.HexColor("#fff8ec")))
            else:
                line += ["", "", "", ""]
            if j == 0:
                line.append("")
        out_rows.append(line)
        style.append(("LINEBELOW", (0, i + 1), (3, i + 1), 0.25, LIGHT))
        style.append(("LINEBELOW", (5, i + 1), (8, i + 1), 0.25, LIGHT))
    cw = [1.75 * inch, 0.48 * inch, 0.5 * inch, 0.47 * inch]
    colw = cw + [0.16 * inch] + cw
    scale = avail / sum(colw)
    return [Table(out_rows, colWidths=[w * scale for w in colw], style=style, repeatRows=1),
            Spacer(1, 3), P("* BOX = second printed price column (box / dozen); never updated when single prices went up. "
                            "Highlighted rows are on an open card.", "note")]

def contacts_table(contacts, avail):
    groups = OrderedDict()
    for c in contacts:
        if c.get("status") == "retired":
            continue
        key = c["organization"] or c["name"]
        groups.setdefault(key, []).append(c)
    data = [[P(f"<font size=6.8 color='#6b7178'>{h}</font>", "small") for h in ("WHO", "", "NUMBER / ADDRESS", "NOTES")]]
    style = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, 0), 0.6, INK),
             ("LEFTPADDING", (0, 0), (-1, -1), 2.5), ("TOPPADDING", (0, 0), (-1, -1), 1.4), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.4)]
    i = 0
    for org, cs in sorted(groups.items(), key=lambda kv: kv[0].lower()):
        first = True
        for c in cs:
            i += 1
            who = esc(org) if first else ""
            if c["name"] and c["name"] != org and c["name"] not in org:
                who = (who + " - " if who else "") + f"<font color='#1d2024'>{esc(c['name'])}</font>"
            fl = f" <font color='#b26a00'><b>[{esc(c['cards'])}]</b></font>" if c.get("status") == "pending" else ""
            data.append([P(("<b>" + who + "</b>") if first else who, "cell"), P(esc(c["kind"]), "note"),
                         P(esc(c["value"]) + fl, "cell"), P(esc(c.get("notes", "")), "note")])
            if first:
                style.append(("LINEABOVE", (0, i), (-1, i), 0.4, RULE))
            first = False
    cw = [2.0 * inch, 0.55 * inch, 2.5 * inch, 2.35 * inch]
    scale = avail / sum(cw)
    return Table(data, colWidths=[w * scale for w in cw], style=style, repeatRows=1)

def index_table(live, pages, avail):
    idx = defaultdict(set)
    for r in live:
        terms = [t.strip().lower() for t in r.get("keywords", "").split(",") if t.strip()]
        if r["category"] != "Pens & Markers":
            terms.append(r["item"].lower())
        for t in terms:
            if len(t) > 1:
                idx[t].add(r["category"])
    entries = []
    for t in sorted(idx):
        cats = sorted(idx[t], key=lambda c: pb.CATEGORY_ORDER.index(c) if c in pb.CATEGORY_ORDER else 99)
        refs = "; ".join(f"{c} {pages.get(c, '')}" for c in cats)
        entries.append(P(f"<b>{esc(t)}</b>  <font color='#6b7178'>{esc(refs)}</font>", S("ix", fontSize=7.2, leading=8.6)))
    ncol = 2
    per = -(-len(entries) // ncol)
    data = []
    for i in range(per):
        data.append([entries[i + j * per] if i + j * per < len(entries) else "" for j in range(ncol)])
    return Table(data, colWidths=[avail / ncol] * ncol, style=[("VALIGN", (0, 0), (-1, -1), "TOP"),
                                                                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                                                                ("TOPPADDING", (0, 0), (-1, -1), 0.8),
                                                                ("BOTTOMPADDING", (0, 0), (-1, -1), 0.8)])

# ============================================================== xlsx
def xlsx(root):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    rows, contacts, cards = pb.load_prices(), pb.load_contacts(), pb.load_cards()
    out = os.path.join(root, "build", "Rainbow_PriceBook.xlsx")
    wb = Workbook()
    def sheet(ws, header, data, widths, money_cols=()):
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1D2024")
        for d in data:
            ws.append(d)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in money_cols:
            for cell in ws[get_column_letter(col)][1:]:
                cell.number_format = '"$"#,##0.00'
    def num(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    cols = ["category", "item", "variant", "size", "color", "sides", "qty_min", "qty_max", "price", "price_max", "unit",
            "min_charge", "status", "cards", "fulfillment", "turnaround", "stock", "brand", "sku", "notes", "id", "source", "updated"]
    def rowvals(r):
        out = []
        for k in cols:
            v = r.get(k, "")
            if k in ("price", "price_max", "min_charge"):
                v = num(v)
            elif k in ("qty_min", "qty_max"):
                v = int(v) if str(v).isdigit() else None
            out.append(v)
        return out
    ws = wb.active; ws.title = "Prices"
    order = {c: i for i, c in enumerate(pb.CATEGORY_ORDER)}
    live = sorted(live_rows(rows), key=lambda r: (order.get(r["category"], 99),) + sort_key(r))
    sheet(ws, [c.replace("_", " ").title() for c in cols], [rowvals(r) for r in live],
          [22, 30, 34, 10, 10, 6, 8, 8, 9, 9, 8, 9, 10, 9, 14, 14, 16, 12, 12, 50, 36, 16, 11], money_cols=(9, 10, 12))
    for i, r in enumerate(live, start=2):
        if r["status"] != "active":
            for c in ws[i]:
                c.fill = PatternFill("solid", fgColor="FFF3DC")
    ws2 = wb.create_sheet("All rows incl. superseded")
    sheet(ws2, [c.replace("_", " ").title() for c in cols], [rowvals(r) for r in rows],
          [22, 30, 34, 10, 10, 6, 8, 8, 9, 9, 8, 9, 10, 9, 14, 14, 16, 12, 12, 50, 36, 16, 11], money_cols=(9, 10, 12))
    ws3 = wb.create_sheet("Phone list")
    sheet(ws3, ["Organization", "Name", "Role", "Kind", "Number / address", "Visibility", "Notes", "Status", "Card"],
          [[c["organization"], c["name"], c["role"], c["kind"], c["value"], c["visibility"], c["notes"], c["status"], c["cards"]]
           for c in contacts if c.get("status") != "retired"], [30, 28, 20, 9, 40, 10, 50, 9, 8])
    ws4 = wb.create_sheet("Open questions")
    sheet(ws4, ["Card", "Status", "Section", "Title", "Proposal"],
          [[c["id"], c["status"], c["section"], c["title"],
            "; ".join(f"{l['key']}: {l['label']} -> {l.get('proposed') or 'no proposal'}" for l in c.get("lines", []))
            or (c["options"][c["proposed_option"]] if c.get("options") and c.get("proposed_option") is not None else "")]
           for c in cards["cards"]], [7, 10, 16, 60, 90])
    ws5 = wb.create_sheet("Change log")
    log = pb.read_csv("changelog.csv")
    sheet(ws5, pb.LOG_FIELDS, [[h.get(k, "") for k in pb.LOG_FIELDS] for h in log], [16, 10, 9, 9, 40, 10, 10, 10, 26, 40])
    wb.save(out)
    return out

# ============================================================== html
def html_page(root, fragment=False):
    rows, contacts, cards = pb.load_prices(), pb.load_contacts(), pb.load_cards()
    keep = ["id", "family", "category", "item", "variant", "brand", "sku", "size", "stock", "color", "sides", "qty_min",
            "qty_max", "unit", "price", "price_max", "min_charge", "fulfillment", "turnaround", "status", "cards", "notes", "keywords"]
    data = [{k: r.get(k, "") for k in keep} for r in live_rows(rows)]
    pub = [{k: c.get(k, "") for k in ("name", "organization", "role", "kind", "value", "notes", "status", "cards")}
           for c in contacts if c.get("visibility") == "public" and c.get("status") != "retired"]
    open_cards = [{"id": c["id"], "title": c.get("status_note") or c["title"], "section": c["section"]} for c in cards["cards"] if c["status"] != "resolved"]
    payload = json.dumps({"generated": dt.date.today().isoformat(), "prices": data, "contacts": pub,
                          "cards": open_cards, "order": pb.CATEGORY_ORDER}, ensure_ascii=False, separators=(",", ":"))
    payload = payload.replace("</", "<\\/")
    tpl = open(os.path.join(root, "templates", "pricebook_page.html"), encoding="utf-8").read()
    body = tpl.replace("/*__DATA__*/null", payload)
    if fragment:
        out = os.path.join(root, "build", "PriceBook_web.html")
        open(out, "w", encoding="utf-8").write(body)
    else:
        out = os.path.join(root, "build", "PriceBook.html")
        open(out, "w", encoding="utf-8").write(
            "<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1, viewport-fit=cover\">\n</head>\n<body>\n"
            + body + "\n</body>\n</html>\n")
    return out

# ============================================================== mirror
def mirror_md(root):
    rows, contacts, cards = pb.load_prices(), pb.load_contacts(), pb.load_cards()
    open_cards = [c for c in cards["cards"] if c["status"] != "resolved"]
    out = os.path.join(root, "build", "pricebook_mirror.md")
    L = [f"# Rainbow price book (mirror, generated {dt.date.today().isoformat()})", "",
         "Read-only copy for price lookups from claude.ai. The master data is on Andrew's Mac at "
         "Rainbow/PriceBook/data/prices.csv (pricebook.py). Price changes made from chat go in "
         "claude/pricebook-pending.md and are applied in the next Cowork session with the Mac connected.",
         "", "Status: PENDING Cxx = proposed from the scans, awaiting the conflict sheet; NO PRICE = no usable price on file. "
         "Sizes are inches, short side first. Prices exclude sales tax.", ""]
    if open_cards:
        L.append(f"Open questions ({len(open_cards)}): " + "; ".join(f"{c['id']} {c['title']}" for c in open_cards))
        L.append("")
    bycat = OrderedDict((c, []) for c in pb.CATEGORY_ORDER)
    for r in live_rows(rows):
        bycat.setdefault(r["category"], []).append(r)
    for cat, rs in bycat.items():
        if not rs:
            continue
        L += [f"## {cat}", ""]
        if CAT_NOTES.get(cat):
            L += [CAT_NOTES[cat], ""]
        L += ["| Item | Details | Size | Qty | Price | Status | id |", "|---|---|---|---|---|---|---|"]
        for r in sorted(rs, key=sort_key):
            det = ", ".join(x for x in (r.get("variant"), r.get("color") if r.get("color") in ("Black ink", "Color ink") else "", r.get("sku")) if x)
            price = (pb.fmt_price(r) + " " + pb.unit_label(r)).strip() if r.get("price") else "no price"
            if r.get("price") and r.get("min_charge"):
                price += f" (min {pb.money(r['min_charge'])})"
            st = pb.flag(r).strip()
            L.append(f"| {r['item']} | {det} | {r.get('size', '')} | {pb.qty_label(r)} | {price} | {st} | {r['id']} |")
        L.append("")
    L += ["## Phone list (vendors and businesses)", "", "| Who | Kind | Number / address | Notes |", "|---|---|---|---|"]
    for c in contacts:
        if c.get("visibility") != "public" or c.get("status") == "retired":
            continue
        who = c["name"] + (f" ({c['organization']})" if c["organization"] and c["organization"] not in c["name"] else "")
        L.append(f"| {who} | {c['kind']} | {c['value']} | {c.get('notes', '')} |")
    open(out, "w", encoding="utf-8").write("\n".join(L) + "\n")
    return out

def build_all(root, draft="auto", only=None):
    outs = []
    want = set(only or ["book", "xlsx", "html", "mirror"])
    if "book" in want:
        outs.append(book_pdf(root, draft))
    if "xlsx" in want:
        outs.append(xlsx(root))
    if "html" in want:
        outs.append(html_page(root)); outs.append(html_page(root, fragment=True))
    if "mirror" in want:
        outs.append(mirror_md(root))
    return outs
