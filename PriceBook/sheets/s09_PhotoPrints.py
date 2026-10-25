"""s09_PhotoPrints.py - Photo Print Pricing (replaces scan_147 p09)."""
import re
import sheetlib as L
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SHEET = {"num": "09", "name": "PhotoPrints", "title": "Photo Print Pricing",
         "source": "scan_147 p09", "covers": ["Photo Prints"]}

ROW_PT = 26
TAB_IN = 6.35         # right-aligned price column, measured from the scan
TEXT_W = 7.3
DARK = L.RGBColor(0x33, 0x33, 0x33)
RED_P = L.RGBColor(0xB5, 0x0A, 0x2E)   # the scan's crimson headline


def _rule(p, sz, color="000000", where="bottom"):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{where}")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "1"); e.set(qn("w:color"), color)
    b.append(e); pPr.append(b)


def _dims(size):
    return tuple(float(x) for x in size.lower().split("x"))


def _fmt(n):
    return str(int(n)) if n == int(n) else str(n)


def build(pb: L.Prices):
    doc = L.new_document(margins_in=(0.5, 0.6, 0.5, 0.6), base_font=L.SANS)

    # live photo sizes, smallest to largest by area (then short side)
    rows = [r for r in pb.live("Photo Prints") if r["size"]]
    rows.sort(key=lambda r: (_dims(r["size"])[0] * _dims(r["size"])[1], _dims(r["size"])))
    for r in rows:
        pb.row(r["id"])

    L.logo_header(doc, logo_width_in=4.4, name_size=19, contact_size=12.5, addr_size=9.5, after=0)

    rule = L.para(doc, before=14, after=0, size=4)
    _rule(rule, 18)

    L.para(doc, "Photo Print Pricing", align="center", before=14, font=L.SANS, size=25,
           bold=True, color=RED_P, line=32)
    lo, hi = rows[0]["size"], rows[-1]["size"]
    L.para(doc, f"All Common Sizes, {lo} through {hi}", align="center", before=2, after=14,
           font=L.SANS, size=14, bold=True, italic=True, color=DARK)

    for i, r in enumerate(rows):
        w, h = _dims(r["size"])
        label = f"{_fmt(w)} x {_fmt(h)}"
        p = L.para(doc, tabs=[(TAB_IN, "right", None)], line=ROW_PT, indent_in=0.0)
        L.add_run(p, label, font=L.SANS, size=20, bold=True)
        L.add_run(p, "\t" + pb.money(r["id"]), font=L.SANS, size=18, bold=True)

    rule2 = L.para(doc, before=4, after=0, size=4)
    _rule(rule2, 12)

    def extra(label, rid, before):
        p = L.para(doc, tabs=[(TAB_IN, "right", None)], before=before, line=24)
        L.add_run(p, label, font=L.SANS, size=16, bold=True)
        L.add_run(p, "\t" + pb.money(rid) + " / photo", font=L.SANS, size=16, bold=True)

    extra("Photo Touch-Up / Color Correction", "photo-touchup", 12)
    extra("Cropping", "photo-cropping", 4)

    rule3 = L.para(doc, before=8, after=0, size=4)
    _rule(rule3, 4, "BBBBBB")
    L.para(doc, "Prices subject to change without notice.", before=7, font=L.SERIF,
           size=11, italic=True, color=DARK)

    L.rev_stamp(doc, pb)
    return doc
