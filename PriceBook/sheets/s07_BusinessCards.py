"""s07_BusinessCards.py - Prices for Printing Business Cards U.S Sizes (replaces scan_147 p07)."""
import sheetlib as L
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SHEET = {"num": "07", "name": "BusinessCards", "title": "Prices for Printing Business Cards U.S Sizes",
         "source": "scan_147 p07", "covers": ["Business Cards"]}

NAVY = L.RGBColor(0x1C, 0x2C, 0x6C)
LEFT_W = 3.9
HEAD_SZ = 32
DESC_SZ = 13.5
NOTE_SZ = 12.8
NOTE_SP = -0.55


def _shade_par(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _notes(doc, lines, before=6):
    for i, ln in enumerate(lines):
        L.para(doc, ln, align="center", before=before if i == 0 else 1, font=L.SANS, size=NOTE_SZ,
               bold=True, char_spacing_pt=NOTE_SP)


def _section(doc, widths, headers, rows, desc, highlight=None, row_h=0.40, head_h=0.40, last_h=None):
    ncol = len(widths)
    t = L.table(doc, 1 + len(rows), ncol, widths_in=widths, align="center")
    mar = t._tbl.tblPr.find(L.qn("w:tblCellMar"))
    for side in ("top", "bottom"):          # zero vertical cell padding so exact heights can't clip
        mar.find(L.qn("w:" + side)).set(L.qn("w:w"), "0")
    left = L.merge(t, 0, 0, len(rows), 0)
    # left block: heading + description (right aligned)
    p = left.paragraphs[0]
    p.paragraph_format.alignment = L.ALIGN["right"]
    L.add_run(p, "Business Cards", font=L.SANS, size=HEAD_SZ, bold=True, underline=True, color=NAVY)
    for ln in desc:
        L.para(left, ln, align="right", font=L.SANS, size=DESC_SZ, line=16.5)
    if highlight:
        q = L.para(left, highlight[0], align="right", before=4, font=L.SANS, size=DESC_SZ, line=16.5,
                   indent_in=0.35)
        _shade_par(q, L.YELLOW_HEX)
        for ln in highlight[1:]:
            q = L.para(left, ln, align="right", font=L.SANS, size=DESC_SZ, line=16.5, indent_in=0.35)
            _shade_par(q, L.YELLOW_HEX)
    left.vertical_alignment = L.WD_CELL_VERTICAL_ALIGNMENT.TOP
    for j, h in enumerate(headers):
        c = t.cell(0, j + 1)
        L.cell_text(c, h, align="center", valign="bottom", font=L.SANS, size=16, bold=True,
                    underline=True, color=NAVY)
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r):
            L.cell_text(t.cell(i, j + 1), v, align="center", valign="top", before=6, font=L.SANS,
                        size=20, bold=True)
    L.set_row_height(t.rows[0], head_h, "exact")
    for i in range(1, len(rows) + 1):
        h = last_h if (last_h and i == len(rows)) else row_h
        L.set_row_height(t.rows[i], h, "exact")
    return t


def build(pb: L.Prices):
    doc = L.new_document(margins_in=(0.3, 0.4, 0.4, 0.4), base_font=L.SANS)
    m = lambda rid, f="price": pb.money(rid, f)

    L.logo_header(doc, logo_width_in=4.3, name_size=19, contact_size=13, addr_size=10, after=4)

    L.para(doc, "Prices for Printing", align="center", font=L.SANS, size=26, bold=True, color=L.RED, line=29)
    L.para(doc, "Business Cards U.S Sizes", align="center", font=L.SANS, size=26, bold=True, color=L.RED,
           line=29)
    L.para(doc, "3.5” x 2” or 2” x 3.5” Any other size requires a special quote.",
           align="center", before=0, font=L.SANS, size=18, bold=True, color=L.RED, line=25)

    # thick black rule
    t = L.table(doc, 1, 1, widths_in=[7.5], align="center")
    L.set_cell_borders(t.cell(0, 0), bottom=(48, "000000"))
    L.set_row_height(t.rows[0], 0.12, "exact")
    L.cell_text(t.cell(0, 0), "", size=2)

    # ---- section 1: outsourced (4over)
    rows1 = [(q, m(f"bc-4over-{q}")) for q in (100, 250, 500, 1000)]
    L.para(doc, "", size=2)
    _section(doc, [LEFT_W, 1.6, 1.9], ["Quantity", "Price"], [(str(q), p) for q, p in rows1],
             ["Printed 1 or 2 sides on 14pt. or 16pt.", "Gloss or Matte Cover or 14pt. Uncoated",
              "These take approx. 1 week for production", "Artwork can be 1 up on a page with",
              "crop marks and bleeds"], row_h=0.44, head_h=0.44, last_h=0.44)
    _notes(doc, ["Prices do not include Artwork and Layout Charges. • Customer can supply artwork.",
                 "Most products take approx. 1 week for production. • E-mail the file to: " + L.SHOP_EMAIL,
                 "The best File format is a High Resolution pdf.  with crop marks and bleeds."], before=6)

    # ---- section 2: in-house
    rows2 = [(str(q), m(f"bc-inhouse-1s-{q}"), m(f"bc-inhouse-2s-{q}")) for q in (100, 200, 250, 500)]
    L.para(doc, "", size=2)
    cut = m("bc-inhouse-bleed-cutting")
    _section(doc, [LEFT_W - 0.1, 1.15, 1.6, 1.05], ["Quantity", "Price 1 sided", "2 sided"], rows2,
             ["Printed “In House” Standard U.S. Size", "On 80# Cover Weight On",
              "any 80# We have in stock", "(2 Sided is not recommended)",
              "Based on Artwork set-up 10 on a page", "with crop marks and bleeds"],
             highlight=["For bleeds, set up is 8 up on a page", "Yielding 20% Less quantity and",
                        f"{cut} extra for cutting."], row_h=0.4767, head_h=0.4767, last_h=0.77)
    _notes(doc, ["Prices do not include Artwork and Layout Charges. • Customer can supply artwork.",
                 "Cards can be either same day or 1 day for production. • E-mail the file to: " + L.SHOP_EMAIL,
                 "The best File format is a High Resolution pdf.  •  8 or 10 up on a page,  with crop marks and bleeds."],
           before=7)

    L.rev_stamp(doc, pb)
    return doc
