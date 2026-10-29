"""s03_ComputerServices.py - Computer usage services and pricing (replaces scan_147 p03)."""
import sheetlib as L
from docx.shared import Pt

SHEET = {"num": "03", "name": "ComputerServices", "title": "COMPUTER USAGE SERVICES AND PRICING",
         "source": "scan_147 p03",
         "covers": ["Computer & File Services", "Scanning", "Copies - B/W", "Copies - Color",
                    "Large Format - B/W", "Large Format - Color", "Large Format - Vellum"]}

S = L.SANS
BODY = 12
HEAD = 13        # red section headings
BLUE = L.BLUE
RED = L.RED


def cp(cell, lines, *, align="center", valign="center", size=BODY, before=0, after=0, line=None):
    """Fill a cell. lines: list of str or (str, fmt-dict). First paragraph is reused."""
    first = True
    for ln in lines:
        txt, fmt = (ln, {}) if isinstance(ln, str) else ln
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.alignment = L.ALIGN[fmt.get("align", align)]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if line:
            p.paragraph_format.line_spacing = line
        if txt:
            L.add_run(p, txt, font=S, size=fmt.get("size", size), bold=fmt.get("bold"),
                      italic=fmt.get("italic"), color=fmt.get("color"))
    if before:
        cell.paragraphs[0].paragraph_format.space_before = Pt(before)
    if after:
        cell.paragraphs[-1].paragraph_format.space_after = Pt(after)
    cell.vertical_alignment = {"top": L.WD_CELL_VERTICAL_ALIGNMENT.TOP,
                               "center": L.WD_CELL_VERTICAL_ALIGNMENT.CENTER}[valign]


def heading(container, text, *, align="left", before=8, after=2, size=HEAD):
    return L.para(container, text, align=align, before=before, after=after, font=S, size=size,
                  color=RED, keep_with_next=True)


def tiny_after(container, pts=2):
    p = L.para(container, "", before=0, after=0)
    p.paragraph_format.line_spacing = Pt(pts)
    r = p.add_run("")
    r.font.size = Pt(1)
    return p


def shrink_trailing(cell):
    """cell.add_table leaves an empty paragraph after the nested table: make it tiny."""
    p = cell.paragraphs[-1]
    if not p.text:
        p.paragraph_format.line_spacing = Pt(1)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("")
        r.font.size = Pt(1)


def grid(t, sz=6):
    """Explicit single-line borders on every cell so inner verticals meet the bottom edge."""
    e = (sz, "000000")
    for row in t.rows:
        for c in row.cells:
            L.set_cell_borders(c, top=e, bottom=e, left=e, right=e)


def build(pb: L.Prices):
    m = lambda rid, f="price", **k: pb.money(rid, f, drop_zero_cents=k.pop("dz", False), **k)
    doc = L.new_document(margins_in=(0.4, 0.5, 0.45, 0.5), base_font=S, base_size=BODY)
    W = [2.3, 2.6, 2.6]

    # ---- title
    L.para(doc, "COMPUTER USAGE SERVICES AND PRICING", align="center", font=S, size=21,
           before=0, after=4)

    # ---- Print from Electronic file
    t = L.table(doc, 1, 3, widths_in=W, borders="all", border_sz=6)
    c = t.rows[0].cells
    cp(c[0], [("Print from Electronic file", {"color": RED, "size": HEAD, "align": "left"}),
              ("", {}), "FLASH DRIVE, EMAIL,", "CAMERA CARD", "Any size copy"], valign="top", before=2)
    cp(c[1], ["1-99 COPIES", "", "NO CHANGES TO FILE",
              f"{m('file-print-no-changes')} ~ {m('file-print-no-changes', 'price_max')} (labor)",
              "PLUS COPY CHARGE"], before=1, after=1)
    cp(c[2], ["1-99 COPIES", "W/CHANGES TO FILE", f"{m('file-print-changes-base')} &",
              f"{m('file-print-changes-hourly', 'min_charge', dz=True)} min/"
              f"{m('file-print-changes-hourly', dz=True)} per hour",
              "PLUS COPY CHARGE"], before=1, after=1)

    # ---- SCAN TO Electronic file
    heading(doc, "SCAN TO Electronic file", before=3)
    t = L.table(doc, 4, 3, widths_in=W, borders="all", border_sz=6)
    r = t.rows
    cp(r[0].cells[0], ["Large format"], align="left")
    cp(r[0].cells[1], [f"{m('scan-lf-bw-first')} first page",
                       f"{m('scan-lf-bw-add')} each following page"], before=1, after=1)
    cp(r[1].cells[0], ["Color Large format"], align="left")
    cp(r[1].cells[1], [f"{m('scan-lf-color-first')} first page",
                       f"{m('scan-lf-color-add')} each following page"], before=1, after=1)
    mc = L.merge(t, 0, 2, 1, 2)
    cp(mc, [f"Large format multiple saves,", f"{m('scan-lf-extra-save')} each extra file"])
    cp(r[2].cells[0], ["8.5x11~11x17", "(document feedable)"], align="left", before=1, after=1)
    mc = L.merge(t, 2, 1, 2, 2)
    cp(mc, [f"{m('scan-doc-first')} first page",
            f"{m('scan-doc-add', cents_style='leading')} each following page"], before=1, after=1)
    cp(r[3].cells[0], ["Photographs", "non-document feedable"], align="left", before=1, after=1)
    mc = L.merge(t, 3, 1, 3, 2)
    cp(mc, [f"{m('scan-photo-base')}",
            f"& {m('scan-photo-hourly')}/Hour until job is completed"], before=1, after=1)

    # ---- OTHER SERVICES
    heading(doc, "OTHER SERVICES:", before=3)
    t = L.table(doc, 1, 2, widths_in=[2.3, 5.2], borders="all", border_sz=6)
    cp(t.rows[0].cells[0], ["Enlarge/Reduce"], align="left", before=1, after=1)
    cp(t.rows[0].cells[1], [f"{m('enlarge-reduce')} each plus copy charge"], before=1, after=1)

    # ---- COPIES (per side)
    L.para(doc, "ALL COPIES ARE CHARGED PER SIDE  (double-sided = 2 sides)", align="center",
           before=4, after=0, font=S, size=12, bold=True)
    outer = L.table(doc, 1, 2, widths_in=[3.9, 3.6], borders=None)
    for cc in outer.rows[0].cells:
        L.set_cell_margins_table  # (no-op reference; cell margins come from table defaults)
    lc, rc = outer.rows[0].cells
    # color
    p = lc.paragraphs[0]
    p.paragraph_format.alignment = L.ALIGN["center"]
    p.paragraph_format.space_before = Pt(2)
    L.add_run(p, "COLOR COPIES", font=S, size=HEAD, color=RED)
    ct = L.table(lc, 4, 3, widths_in=[1.2, 1.3, 1.3], borders="all", border_sz=6)
    shrink_trailing(lc)
    for i, h in enumerate(["Size", "1-20 sides", "21 & up sides"]):
        cp(ct.rows[0].cells[i], [(h, {"bold": True, "italic": True})])
    for k, (sz, a, b) in enumerate([("8.5x11", "color-8.5x11-side-1-20", "color-8.5x11-side-21up"),
                                    ("8.5x14", "color-8.5x14-side-1-20", "color-8.5x14-side-21up"),
                                    ("11x17", "color-11x17-side-1-20", "color-11x17-side-21up")], start=1):
        cp(ct.rows[k].cells[0], [sz], before=1, after=1)
        cp(ct.rows[k].cells[1], [m(a, dollar=False, cents_style="leading")])
        cp(ct.rows[k].cells[2], [m(b, dollar=False, cents_style="leading")])
    L.para(lc, f"Cost per side \u00b7 100 lb cardstock +{m('stock-100lb-cover-8.5x11')}/sheet",
           align="center", before=2, font=S, size=10, italic=True)
    grid(ct)
    # b/w
    p = rc.paragraphs[0]
    p.paragraph_format.alignment = L.ALIGN["center"]
    p.paragraph_format.space_before = Pt(2)
    L.add_run(p, "BLACK/WHITE COPIES", font=S, size=HEAD, color=RED)
    bt = L.table(rc, 4, 2, widths_in=[1.7, 1.7], borders="all", border_sz=6)
    shrink_trailing(rc)
    cp(bt.rows[0].cells[0], [("Size", {"bold": True, "italic": True})])
    cp(bt.rows[0].cells[1], [("Cost per side", {"bold": True, "italic": True})])
    for i, (sz, rid) in enumerate([("8.5x11", "bw-8.5x11-side"), ("8.5x14", "bw-8.5x14-side"),
                                   ("11x17", "bw-11x17-side")], start=1):
        cp(bt.rows[i].cells[0], [sz], before=1, after=1)
        cp(bt.rows[i].cells[1], [m(rid, dollar=False, cents_style="leading")])
    L.para(rc, "Cost per side - no quantity price breaks", align="center", before=2, font=S,
           size=10, italic=True)

    grid(bt)
    L.para(doc, f"Hole punching or stapling: add {m('copy-addon-punch-staple', dollar=False, cents_style='leading')} per page to the copy charge",
           align="center", before=1, after=0, font=S, size=10, bold=True)

    # ---- Large format
    outer = L.table(doc, 1, 2, widths_in=[3.9, 3.6], borders=None)
    lc, rc = outer.rows[0].cells
    p = lc.paragraphs[0]
    p.paragraph_format.alignment = L.ALIGN["center"]
    p.paragraph_format.space_before = Pt(3)
    L.add_run(p, "Large Format Color", font=S, size=HEAD, color=RED)
    ft = L.table(lc, 7, 3, widths_in=[1.5, 1.25, 1.15], borders="all", border_sz=6)
    shrink_trailing(lc)
    g = L.merge(ft, 0, 0, 3, 0)
    cp(g, ["Color blueprints/plans", "(flat price each)"])
    for i, (sz, rid) in enumerate([("18x24", "lf-color-18x24"), ("24x36", "lf-color-24x36"),
                                   ("30x42", "lf-color-30x42"), ("36x48", "lf-color-36x48")]):
        cp(ft.rows[i].cells[1], [sz], before=1, after=1)
        cp(ft.rows[i].cells[2], [m(rid)])
    g = L.merge(ft, 4, 0, 4, 1)
    cp(g, ["Bond, posters/graphics, any size"], before=1, after=1)
    cp(ft.rows[4].cells[2], [f"{m('lf-color-bond-sqft', dz=True)} per sq ft"])
    g = L.merge(ft, 5, 0, 5, 1)
    cp(g, ["Photo paper, custom sizes only"], before=1, after=1)
    cp(ft.rows[5].cells[2], [f"{m('lf-color-photo-sqft', dz=True)} per sq ft"])
    g = L.merge(ft, 6, 0, 6, 1)
    cp(g, ["Transparency, any size"], before=1, after=1)
    cp(ft.rows[6].cells[2], [f"{m('lf-color-transparency-sqft', dz=True)} per sq ft"])

    p = rc.paragraphs[0]
    p.paragraph_format.alignment = L.ALIGN["center"]
    p.paragraph_format.space_before = Pt(3)
    L.add_run(p, "Large Format B/W", font=S, size=HEAD, color=RED)
    L.add_run(p, "   each", font=S, size=HEAD - 1, italic=True)
    bt = L.table(rc, 4, 2, widths_in=[1.7, 1.7], borders="all", border_sz=6)
    shrink_trailing(rc)
    for i, (sz, rid) in enumerate([("18x24", "lf-bw-18x24-walkin"), ("24x36", "lf-bw-24x36-walkin"),
                                   ("30x42", "lf-bw-30x42-walkin"), ("36x48", "lf-bw-36x48-walkin")]):
        cp(bt.rows[i].cells[0], [sz], before=1, after=1)
        cp(bt.rows[i].cells[1], [m(rid)])

    # blue lines under the B/W table
    L.para(rc, f"{m('lf-bw-other-sqft')} per sq ft otherwise on B/W large format", before=6,
           font=S, size=BODY - 1.5, color=BLUE, indent_in=0.05)
    L.para(rc, f"Copy onto vellum:  18x24 ... {m('vellum-18x24')}   24x36 ... {m('vellum-24x36')}",
           before=3, font=S, size=BODY - 1.5, color=BLUE, indent_in=0.05)

    # ---- Self serve
    tiny_after(doc, 2)
    t = L.table(doc, 2, 2, widths_in=[3.0, 4.5], borders="all", border_sz=6)
    cp(t.rows[0].cells[0], [("Self Serve Computer", {"color": RED, "size": HEAD})], before=1, after=1)
    cp(t.rows[0].cells[1], [f"{m('self-serve-computer', cents_style='cents')} per minute"])
    cp(t.rows[1].cells[0], [("Self Serve Computer copies", {"color": RED, "size": HEAD})], before=1, after=1)
    cp(t.rows[1].cells[1], [f"{m('self-serve-copies', cents_style='cents')} per side, black/white",
                            ("(double-sided = 2 sides)", {"size": 10})])

    L.rev_stamp(doc, pb)
    return doc
