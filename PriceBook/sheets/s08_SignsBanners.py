"""s08_SignsBanners.py - Vinyl banners, foam core, sandwich boards, coroplast, car magnets (replaces scan_147 p08).

2026-10-08 (scan_200): 10 mm coroplast dropped; 18x24 / 24x36 each-prices added under foam core and 4 mm coroplast.
"""
import sheetlib as L
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SHEET = {"num": "08", "name": "SignsBanners",
         "title": "Prices for Printing Vinyl Banners, Foam Core, Aluminum Sandwich Boards, "
                  "Coroplast Signs and Car Magnets from 4-Over",
         "source": "scan_147 p08",
         "covers": ["Signs & Banners (4over)"]}

LEFT_IN = 3.95      # product-name column (right aligned)
GAP_IN = 0.30
RIGHT_IN = 3.45
NAME_PT = 29
DESC_PT = 12.5
PRICE_PT = 17


def _bottom_rule(p, sz=36):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "1")
    e.set(qn("w:color"), "000000")
    b.append(e)
    pPr.append(b)


def build(pb: L.Prices):
    doc = L.new_document(margins_in=(0.45, 0.4, 0.5, 0.4), base_font=L.SANS)
    m = lambda rid: pb.money(rid) + " Per Sq. Ft."

    L.logo_header(doc, logo_width_in=4.0, name_size=19, contact_size=12, addr_size=8.5, after=8)

    p = L.para(doc, align="center", font=L.SANS, size=22, bold=True, color=L.RED, line=24, after=6,
               text="Prices for Printing Vinyl Banners, Foam Core,")
    p.add_run().add_break()
    L.add_run(p, "Aluminum Sandwich Boards, Coroplast Signs", font=L.SANS, size=22, bold=True, color=L.RED)
    p.add_run().add_break()
    L.add_run(p, "and Car Magnets from 4-Over", font=L.SANS, size=22, bold=True, color=L.RED)
    p.paragraph_format.left_indent = Pt(10)
    p.paragraph_format.right_indent = Pt(10)
    rule = L.para(doc, "", after=10)
    _bottom_rule(rule)
    rule.paragraph_format.space_after = Pt(6)

    # Each item is one or more table rows so every price sits exactly level with its label.
    #   ("name", name_pt, price_id)       product line: big name, price per sq ft
    #   ("desc", text, price_id or None)  description line, optional price on the same line
    #   ("size", price_id)                standard-size line: 18" x 24"  ->  $xx.xx each
    SIZE_PT = 15
    items = [
        [("name", NAME_PT, "banner-13oz"),
         ("desc", "Waterproof, durable 13 oz scrim vinyl", None),
         ("desc", "Economical, light 8 oz mesh vinyl", "banner-8oz-mesh")],
        [("name", NAME_PT, "foam-core"),
         ("size", "foam-core-18x24"),
         ("size", "foam-core-24x36")],
        [("name", NAME_PT, "coroplast-4mm"),
         ("desc", "Waterproof, lightweight, traditional yard signs", None),
         ("size", "coroplast-4mm-18x24"),
         ("size", "coroplast-4mm-24x36")],
        [("name", 19, "aluminum-sandwich"),
         ("desc", "Durable, lightweight aluminum sandwich board", None),
         ("desc", "signs are composed of two aluminum sheets with", None),
         ("desc", "polyethylene core, resulting in a total thickness", None),
         ("desc", 'of 1/8" (3 mm)', None)],
        [("name", 24, "car-magnet")],
    ]
    names = {"banner-13oz": "Banners", "foam-core": "Foam Core", "coroplast-4mm": "4 mm Coroplast",
             "aluminum-sandwich": "Aluminum Sandwich Boards", "car-magnet": "Car Magnets"}

    def size_label(rid):
        w, h = pb.row(rid)["size"].lower().split("x")
        return f'{w}" x {h}"'

    flat = [(i, ln) for i, it in enumerate(items) for ln in it]
    t = L.table(doc, len(flat), 3, widths_in=[LEFT_IN, GAP_IN, RIGHT_IN], align="center")
    desc_line = round(DESC_PT * 1.25)
    size_line = round(SIZE_PT * 1.25)
    for r, (i, ln) in enumerate(flat):
        lc, rc = t.cell(r, 0), t.cell(r, 2)
        kind = ln[0]
        if kind == "name":
            _, npt, rid = ln
            nl = round(npt * 1.15)
            before = 0 if r == 0 else 6
            L.cell_text(lc, names[rid], align="right", valign="bottom", font=L.SANS, size=npt, bold=True,
                        line=nl, before=before)
            L.cell_text(rc, m(rid), align="left", valign="bottom", font=L.SANS, size=PRICE_PT, bold=True,
                        line=nl, before=before)
        elif kind == "desc":
            _, text, rid = ln
            L.cell_text(lc, text, align="right", valign="bottom", font=L.SANS, size=DESC_PT, line=desc_line)
            L.cell_text(rc, m(rid) if rid else "", align="left", valign="bottom", font=L.SANS, size=PRICE_PT,
                        bold=True, line=desc_line)
        else:  # size
            _, rid = ln
            L.cell_text(lc, size_label(rid), align="right", valign="bottom", font=L.SANS, size=SIZE_PT,
                        bold=True, line=size_line)
            L.cell_text(rc, pb.money(rid) + " each", align="left", valign="bottom", font=L.SANS,
                        size=SIZE_PT, bold=True, line=size_line)
    L.no_split_rows(t)

    L.para(doc, "Prices do not include Artwork and Layout Charges. Most products take approx. 1 week for production.",
           align="left", before=14, indent_in=0.05, font=L.SANS, size=11.5)

    L.rev_stamp(doc, pb)
    return doc
