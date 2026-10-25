"""sheetlib.py - shared helpers for Rainbow's laminated counter price sheets.

Every counter sheet is a small Python module (sNN_name.py) that builds a python-docx
Document from the master price list (data/prices.csv).  Prices are NEVER typed into a
sheet module: every dollar figure comes from Prices.money(<row id>) so that
`pricebook.py set ...` + `python3 sheets/build_sheets.py` updates every sheet.

Layout of the kit (inside PriceBook/):
    sheets/sheetlib.py        this file
    sheets/build_sheets.py    builds all sheets (or the ones named) -> .docx + .pdf
    sheets/audit_sheets.py    deterministic checks (prices on the PDF == prices.csv, coverage)
    sheets/sNN_*.py           one module per sheet
    sheets/assets/            logo etc.
    Counter Sheets/           output: Rainbow_Sheet_NN_<Name>.docx / .pdf (+ .manifest.json)

Fonts: use only "Times New Roman" and "Arial" (regular/bold/italic).  On the Mac Word
uses the real fonts; LibreOffice on Linux substitutes the metric-identical Liberation
Serif / Liberation Sans, so line breaks and page fit are the same everywhere.
"""
from __future__ import annotations

import csv
import datetime as _dt
import json
import os
import re
import shutil
import subprocess
import tempfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Emu

HERE = Path(__file__).resolve().parent
PRICEBOOK = HERE.parent
DEFAULT_CSV = Path(os.environ.get("RAINBOW_PRICES_CSV", PRICEBOOK / "data" / "prices.csv"))
ASSETS = HERE / "assets"
LOGO = ASSETS / "rainbow_logo.png"          # crop from scan 147; replace with original art later
OUT_DIR = Path(os.environ.get("RAINBOW_SHEETS_OUT", PRICEBOOK / "Counter Sheets"))

# Shop identity exactly as printed on the logo-header sheets (p07/p08/p09/p11).
SHOP_NAME = "PRINTING & OFFICE SUPPLY"
SHOP_PHONE_LINE = "(775) 831-2276 • FAX (775) 831-5535"
SHOP_ADDR_LINE = "917 TAHOE BLVD., SUITE 103 • BOX 5488 • INCLINE VILLAGE, NV 89451"
SHOP_EMAIL = "RainbowIncline@gmail.com"

SERIF = "Times New Roman"
SANS = "Arial"

# Colours sampled from the scans (approximate; keep consistent across sheets)
RED = RGBColor(0xD0, 0x10, 0x10)       # "Prices for Printing" red headlines
BLUE = RGBColor(0x1F, 0x3F, 0x9F)      # "Business Cards" blue headings
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x66, 0x66, 0x66)
YELLOW_HEX = "FFFF00"                  # highlighter boxes
SHEET_BLUE_HEX = "29B6E6"              # blue card stock (p12/p13)

LIVE = ("active", "pending")


class SheetDataError(Exception):
    pass


# --------------------------------------------------------------------------- prices
class Prices:
    """Read-only view of prices.csv that records which rows a sheet used."""

    def __init__(self, path: Path | str = DEFAULT_CSV):
        self.path = Path(path)
        with open(self.path, newline="", encoding="utf-8") as f:
            self.rows = list(csv.DictReader(f))
        self.by_id = {r["id"]: r for r in self.rows}
        self.used: list[str] = []
        self.excluded: dict[str, str] = {}

    # -- lookups ---------------------------------------------------------
    def row(self, rid: str) -> dict:
        r = self.by_id.get(rid)
        if r is None:
            raise SheetDataError(f"price id not found: {rid}")
        if r["status"] not in LIVE:
            raise SheetDataError(f"{rid} is {r['status']} - not allowed on a counter sheet")
        if r["status"] == "active" and not r["price"].strip() and "Estimate" not in r.get("notes", ""):
            # active rows with no price are 'estimate' style rows (e.g. labels-artwork)
            pass
        if rid not in self.used:
            self.used.append(rid)
        return r

    def has(self, rid: str) -> bool:
        r = self.by_id.get(rid)
        return bool(r and r["status"] in LIVE)

    def value(self, rid: str, field: str = "price") -> Decimal:
        r = self.row(rid)
        v = r.get(field, "").strip()
        if not v:
            raise SheetDataError(f"{rid} has no {field}")
        return Decimal(v)

    def money(self, rid: str, field: str = "price", *, dollar: bool = True,
              cents_style: str = "decimal", drop_zero_cents: bool = False) -> str:
        """Format a price.  cents_style: 'decimal' -> $0.25 ; 'leading' -> $.25 ;
        'cents' -> '25 cents' (only for < $1).  drop_zero_cents: $20 instead of $20.00."""
        d = self.value(rid, field).quantize(Decimal("0.01"), ROUND_HALF_UP)
        return fmt_money(d, dollar=dollar, cents_style=cents_style, drop_zero_cents=drop_zero_cents)

    def live(self, category: str | None = None, *, pred=None) -> list[dict]:
        """All active+pending rows (optionally one category), in CSV order."""
        out = []
        for r in self.rows:
            if r["status"] not in LIVE:
                continue
            if category and r["category"] != category:
                continue
            if pred and not pred(r):
                continue
            out.append(r)
        return out

    def exclude(self, rid: str, reason: str):
        """Declare that a live row in a covered category is intentionally NOT printed."""
        self.excluded[rid] = reason

    def latest_update(self) -> str:
        dates = [self.by_id[i].get("updated", "") for i in self.used if self.by_id[i].get("updated")]
        return max(dates) if dates else _dt.date.today().isoformat()


def fmt_money(d: Decimal, *, dollar=True, cents_style="decimal", drop_zero_cents=False) -> str:
    d = Decimal(d).quantize(Decimal("0.01"), ROUND_HALF_UP)
    if cents_style == "cents" and d < 1:
        return f"{int(d * 100)} cents"
    s = f"{d:,.2f}"
    if drop_zero_cents and s.endswith(".00"):
        s = s[:-3]
    if cents_style == "leading" and s.startswith("0."):
        s = s[1:]
    return ("$" if dollar else "") + s


# --------------------------------------------------------------------------- document
def new_document(*, margins_in=(0.5, 0.5, 0.5, 0.5), landscape=False,
                 base_font=SERIF, base_size=12) -> Document:
    """Letter-size document. margins_in = (top, right, bottom, left)."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    t, r, b, l = margins_in
    sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = (
        Inches(t), Inches(r), Inches(b), Inches(l))
    sec.header_distance = Inches(0.2)
    sec.footer_distance = Inches(0.2)
    st = doc.styles["Normal"]
    st.font.name = base_font
    st.font.size = Pt(base_size)
    _set_font_all(st.element.get_or_add_rPr(), base_font)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    # remove the default empty first paragraph's spacing quirks
    return doc


def _set_font_all(rPr, name):
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), name)


def style_run(run, *, font=None, size=None, bold=None, italic=None, underline=None,
              color: RGBColor | None = None, caps=None, highlight_hex: str | None = None,
              char_spacing_pt: float | None = None):
    if font:
        run.font.name = font
        _set_font_all(run._element.get_or_add_rPr(), font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if underline is not None:
        run.font.underline = underline
    if color is not None:
        run.font.color.rgb = color
    if caps is not None:
        run.font.all_caps = caps
    if highlight_hex:
        rPr = run._element.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), highlight_hex)
        rPr.append(shd)
    if char_spacing_pt is not None:
        rPr = run._element.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(char_spacing_pt * 20)))
        rPr.append(sp)
    return run


ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def para(container, text: str = "", *, align="left", before=0, after=0, line=None,
         keep_with_next=False, indent_in=None, first_line_in=None, tabs=None, **runfmt):
    """Add a paragraph to a Document or table cell.  runfmt -> style_run kwargs.
    Returns the paragraph; add more runs with add_run(p, text, **fmt).
    tabs: list of (position_inches, 'left'|'right'|'center'|'decimal', leader|None)."""
    if hasattr(container, "add_paragraph"):
        p = container.add_paragraph()
    else:
        raise TypeError("container must support add_paragraph")
    pf = p.paragraph_format
    pf.alignment = ALIGN[align]
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        if isinstance(line, (int, float)) and line > 3:
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(line)
        else:
            pf.line_spacing = line
    if keep_with_next:
        pf.keep_with_next = True
    if indent_in is not None:
        pf.left_indent = Inches(indent_in)
    if first_line_in is not None:
        pf.first_line_indent = Inches(first_line_in)
    if tabs:
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        amap = {"left": WD_TAB_ALIGNMENT.LEFT, "right": WD_TAB_ALIGNMENT.RIGHT,
                "center": WD_TAB_ALIGNMENT.CENTER, "decimal": WD_TAB_ALIGNMENT.DECIMAL}
        lmap = {None: WD_TAB_LEADER.SPACES, "dots": WD_TAB_LEADER.DOTS}
        for pos, al, lead in tabs:
            pf.tab_stops.add_tab_stop(Inches(pos), amap[al], lmap[lead])
    if text:
        add_run(p, text, **runfmt)
    return p


def add_run(p, text, **fmt):
    return style_run(p.add_run(text), **fmt)


def first_cell_para(cell):
    """A fresh table cell already holds one empty paragraph: return it (use instead of para())."""
    return cell.paragraphs[0]


def cell_text(cell, text, *, align="left", before=0, after=0, line=None, valign="center", **runfmt):
    """Write text into a cell's first paragraph (replacing it). Newlines -> separate paragraphs."""
    lines = str(text).split("\n")
    p = cell.paragraphs[0]
    for i, ln in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = ALIGN[align]
        pf.space_before = Pt(before if i == 0 else 0)
        pf.space_after = Pt(after if i == len(lines) - 1 else 0)
        if line is not None:
            if line > 3:
                pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                pf.line_spacing = Pt(line)
            else:
                pf.line_spacing = line
        if ln:
            add_run(p, ln, **runfmt)
    cell.vertical_alignment = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                               "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                               "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}[valign]
    return p


# --------------------------------------------------------------------------- tables
def table(doc_or_cell, rows: int, cols: int, *, widths_in=None, align="center",
          borders: str | None = None, border_sz=4, border_color="000000", autofit=False):
    """Create a table.  borders: None/'none' | 'all' | 'outer' | 'grid-inner' .
    border_sz in eighths of a point (4 = 0.5pt, 8 = 1pt, 12 = 1.5pt, 18 = 2.25pt)."""
    t = doc_or_cell.add_table(rows=rows, cols=cols)
    t.alignment = {"left": WD_TABLE_ALIGNMENT.LEFT, "center": WD_TABLE_ALIGNMENT.CENTER,
                   "right": WD_TABLE_ALIGNMENT.RIGHT}[align]
    t.autofit = autofit
    _tbl_layout_fixed(t)
    if widths_in:
        set_col_widths(t, widths_in)
    set_table_borders(t, borders or "none", sz=border_sz, color=border_color)
    set_cell_margins_table(t, 0.04, 0.06)
    return t


def _tbl_layout_fixed(t):
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)


def set_col_widths(t, widths_in):
    tblGrid = t._tbl.tblGrid
    for i, gc in enumerate(tblGrid.findall(qn("w:gridCol"))):
        if i < len(widths_in):
            gc.set(qn("w:w"), str(int(widths_in[i] * 1440)))
    for row in t.rows:
        for i, c in enumerate(row.cells):
            if i < len(widths_in):
                c.width = Inches(widths_in[i])
    # total table width
    tblPr = t._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(int(sum(widths_in) * 1440)))


def set_table_borders(t, mode="all", *, sz=4, color="000000", val="single"):
    tblPr = t._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    edges = {"none": {}, "all": dict.fromkeys(["top", "left", "bottom", "right", "insideH", "insideV"], True),
             "outer": dict.fromkeys(["top", "left", "bottom", "right"], True),
             "grid-inner": dict.fromkeys(["insideH", "insideV"], True),
             "rows": dict.fromkeys(["top", "bottom", "insideH"], True)}[mode]
    for e in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{e}")
        if edges.get(e):
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        el.set(qn("w:space"), "0")
        b.append(el)
    tblPr.append(b)


def set_cell_borders(cell, **edges):
    """set_cell_borders(c, top=(sz,'000000'), bottom=(8,'000000'), left=None ...); None -> nil."""
    tcPr = cell._tc.get_or_add_tcPr()
    tb = tcPr.find(qn("w:tcBorders"))
    if tb is None:
        tb = OxmlElement("w:tcBorders")
        tcPr.append(tb)
    for edge, spec in edges.items():
        el = tb.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tb.append(el)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            sz, color = spec[0], spec[1]
            el.set(qn("w:val"), spec[2] if len(spec) > 2 else "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), color)
            el.set(qn("w:space"), "0")


def shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins_table(t, top_bottom_in=0.03, left_right_in=0.06):
    tblPr = t._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top_bottom_in), ("bottom", top_bottom_in),
                    ("left", left_right_in), ("right", left_right_in)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(v * 1440)))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblPr.append(m)


def set_row_height(row, inches, rule="atLeast"):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(inches * 1440)))
    h.set(qn("w:hRule"), rule)   # 'atLeast' | 'exact'
    trPr.append(h)


def no_split_rows(t):
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)


def merge(t, r1, c1, r2, c2):
    return t.cell(r1, c1).merge(t.cell(r2, c2))


# --------------------------------------------------------------------------- header / images
def picture(container_para, path: Path | str, width_in: float):
    run = container_para.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    return run


def logo_header(doc, *, logo_width_in=4.6, name_size=17, contact_size=10.5, addr_size=6.5,
                after=6, show_addr=True):
    """The RAINBOW logo + PRINTING & OFFICE SUPPLY block used on the newer sheets."""
    p = para(doc, align="center")
    picture(p, LOGO, logo_width_in)
    para(doc, SHOP_NAME, align="center", font=SANS, size=name_size, bold=True, line=1.0)
    para(doc, SHOP_PHONE_LINE, align="center", font=SANS, size=contact_size, bold=True)
    if show_addr:
        para(doc, SHOP_ADDR_LINE, align="center", font=SANS, size=addr_size, bold=True, after=after)


def rev_stamp(doc, pb: Prices, *, size=6, align="right", text=None):
    """Tiny revision line in the page footer ('Rev. YYYY-MM-DD').  Not a price."""
    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.alignment = ALIGN[align]
    add_run(fp, text or f"Rev. {pb.latest_update()}", font=SANS, size=size, color=GRAY)


# --------------------------------------------------------------------------- output
def soffice_bin():
    for b in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              "C:/Program Files/LibreOffice/program/soffice.exe"):
        if shutil.which(b) or Path(b).exists():
            return shutil.which(b) or b
    raise RuntimeError("LibreOffice (soffice) not found - needed to render the PDF")


def to_pdf(docx_path: Path) -> Path:
    """Render docx -> pdf beside it.  LibreOffice works in a private temp dir (its lock and
    temp files never land in the shop folder); only the finished PDF is copied over."""
    docx_path = Path(docx_path)
    prof = tempfile.mkdtemp(prefix="lo_prof_")
    work = tempfile.mkdtemp(prefix="lo_work_")
    try:
        src = Path(work) / docx_path.name
        shutil.copyfile(docx_path, src)
        subprocess.run([soffice_bin(), f"-env:UserInstallation={Path(prof).as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", work, str(src)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        made = src.with_suffix(".pdf")
        if not made.exists():
            raise RuntimeError(f"PDF not produced for {docx_path}")
        pdf = docx_path.with_suffix(".pdf")
        shutil.copyfile(made, pdf)
    finally:
        shutil.rmtree(prof, ignore_errors=True)
        shutil.rmtree(work, ignore_errors=True)
    return pdf


def bottom_slack_in(pdf: Path, docx_path: Path, page: int = 1) -> float:
    """Inches of free space between the last line of body text on `page` and the bottom margin.
    Different LibreOffice/Word versions reflow slightly; keep >= MIN_SLACK_IN so the sheet stays one page."""
    d = Document(str(docx_path))
    s = d.sections[0]
    limit = s.page_height.pt - s.bottom_margin.pt
    x = subprocess.run(["pdftotext", "-bbox", "-f", str(page), "-l", str(page), str(pdf), "-"],
                       capture_output=True, text=True).stdout
    ys = [float(b) for a, b in re.findall(r'yMin="([\d.]+)" xMax="[\d.]+" yMax="([\d.]+)"', x)
          if float(a) < limit]
    return (limit - max(ys)) / 72 if ys else 99.0


MIN_SLACK_IN = 0.35


def pdf_pages(pdf: Path) -> int:
    out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    m = re.search(r"Pages:\s+(\d+)", out)
    return int(m.group(1)) if m else -1


def save_sheet(doc, sheet: dict, pb: Prices, out_dir: Path | None = None, *, pdf=True) -> dict:
    """Save <out>/Rainbow_Sheet_<NN>_<Name>.docx (+ .pdf + .manifest.json). Returns the manifest."""
    out_dir = Path(out_dir or OUT_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"Rainbow_Sheet_{sheet['num']}_{sheet['name']}"
    docx_path = out_dir / f"{stem}.docx"
    doc.core_properties.title = sheet.get("title", stem)
    doc.core_properties.author = "Rainbow Printing & Office Supply"
    doc.core_properties.comments = f"Generated from PriceBook/data/prices.csv - rebuild with sheets/build_sheets.py"
    doc.save(docx_path)
    man = {
        "sheet": sheet, "docx": docx_path.name, "generated": _dt.datetime.now().isoformat(timespec="seconds"),
        "prices_csv": str(pb.path), "used_ids": pb.used,
        "pending_ids": [i for i in pb.used if pb.by_id[i]["status"] == "pending"],
        "excluded": pb.excluded, "rev": pb.latest_update(),
    }
    if pdf:
        p = to_pdf(docx_path)
        man["pdf"] = p.name
        man["pages"] = pdf_pages(p)
    (out_dir / f"{stem}.manifest.json").write_text(json.dumps(man, indent=1), encoding="utf-8")
    return man
