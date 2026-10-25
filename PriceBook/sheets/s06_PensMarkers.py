"""s06_PensMarkers.py - Pens & Markers counter sheet (replaces scan_147 p06).

Two-page quick reference, printed duplex and laminated:
  * strict A-Z order, a coloured letter bar at the start of every letter (rainbow cycle);
    a "(continued)" bar heads any column that starts part-way through a letter or brand
  * brands with GROUP_MIN or more lines get a brand sub-header; the brand word is then
    dropped from each line (BIC / EXPO / PAPERMATE / PENTEL / PILOT / SHARPIE / UNIBALL ...)
  * product-type tints: highlighters yellow, pencils grey, dry/wet-erase blue (legend on top)
  * guide words + letter strip at the top of each side; bold prices, hairline row rules
Layout is computed from the live rows, so adding/retiring a pen just reflows the sheet.
No prices in this file - every figure comes from data/prices.csv via pb.money().
"""
import math
import re
from pathlib import Path

import sheetlib as L
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SHEET = {"num": "06", "name": "PensMarkers", "title": "PENS & MARKERS",
         "source": "scan_147 p06", "covers": ["Pens & Markers"], "pages": 2}

# ---- geometry (inches) ------------------------------------------------------------
MARGINS = (0.40, 0.40, 0.45, 0.40)            # top, right, bottom, left
NAME_W, PRICE_W, GAP_W = 3.14, 0.62, 0.18      # 2*(3.14+0.62)+0.18 = 7.70 in
CELL_LR = 0.06
HEADER_IN = 0.95                               # title + strip + spacer, measured generously
ROW_MAX_IN, ROW_MIN_IN = 0.31, 0.24
SLACK_IN = 0.45                                # keep the audit's 0.35 in with room to spare
INDENT_IN = 0.14                               # lines under a brand header

# ---- type ---------------------------------------------------------------------------
NAME_PT, NAME_MIN_PT = 13, 10
SKU_PT, SKU_MIN_PT = 8.5, 7.5
PRICE_PT = 13.5
LETTER_PT = 15
BRAND_PT = 11
GROUP_MIN = 3                                  # brand sub-header when a brand has >= 3 lines

# ---- colour -------------------------------------------------------------------------
# Letter bars cycle through the rainbow.  Yellow is left out on purpose: it is the
# highlighter tint, and white type on yellow doesn't read.
LETTER_COLORS = ["C62828", "E65100", "2E7D32", "00838F", "1565C0", "3949AB", "6A1B9A"]
TINTS = {"hl": ("FFF59D", "Highlighter"), "pencil": ("E0E0E0", "Pencil"),
         "erase": ("CDE7FA", "Dry / wet erase")}
RULE = (4, "C8C8C8")
SKU_GRAY = RGBColor(0x70, 0x70, 0x70)
STRIP_OFF = RGBColor(0xC4, 0xC4, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Product type per row id.  Rows not listed fall back to keywords in the item name.
TYPE = {
    "pen-bic-briteliner-bl11yw": "hl", "pen-frixion-highlighter": "hl",
    "pen-hi-liter-large-24000": "hl", "pen-hi-liter-pen-style-23591": "hl",
    "pen-pilot-spotliter-45011": "hl", "pen-pilot-spotliter-supreme-16008": "hl",
    "pen-sharpie-accent-liquid-24425": "hl", "pen-sharpie-liquid-highlighter": "hl",
    "pen-sharpie-pocket-accent-27025": "hl", "pen-sharpie-smear-guard-larger-25025": "hl",
    "pen-sharpie-smear-guard-rt-highlighter-28025": "hl",
    "pen-sharpie-smear-guard-narrow-27025": "hl", "pen-zebra-bible-highlighter": "hl",
    "pen-zebra-literite-81150": "hl", "pen-zebra-zazzle-brights-71050": "hl",
    "pen-bic-velocity-pencil-mv711bk": "pencil", "pen-client-pencil-al905a": "pencil",
    "pen-energize-pencil-pl77a": "pencil", "pen-icy-pentel-al25ta": "pencil",
    "pen-pentel-e-sharp": "pencil", "pen-pentel-twist-erase-express": "pencil",
    "pen-pentel-twist-erase-gt": "pencil", "pen-zebra-z-grip-pencil": "pencil",
    "pen-expo-dry-erase-83001": "erase", "pen-expo-finepoint-84001": "erase",
    "pen-expo-w-eraser-80788": "erase", "pen-liquid-expo-89081": "erase",
    "pen-vis-a-vis-16001": "erase",
}
NAME_OVERRIDE = {"pen-pentel-energel-margin": "Pentel Energel"}


def _type(r):
    if r["id"] in TYPE:
        return TYPE[r["id"]]
    n = r["item"].lower()
    if re.search(r"highlight|hi-?liter|spotliter|accent", n):
        return "hl"
    if "pencil" in n:
        return "pencil"
    if re.search(r"dry.?erase|wet.?erase|\bexpo\b", n):
        return "erase"
    return None


# ---- text measuring (for auto-shrinking long names) ---------------------------------
_FONT_PATHS = [
    str(Path(__file__).resolve().parent / "assets" / "LiberationSans-Regular.ttf"),
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf", "C:/Windows/Fonts/arial.ttf",
]
_fonts = {}


def _text_w_in(text, pt):
    try:
        from PIL import ImageFont
        key = round(pt * 4)
        if key not in _fonts:
            path = next(p for p in _FONT_PATHS if Path(p).exists())
            _fonts[key] = ImageFont.truetype(path, size=100)
        return _fonts[key].getlength(text) / 100 * pt / 72
    except Exception:  # noqa: BLE001 - no PIL / no font: conservative estimate
        return len(text) * 0.56 * pt / 72


# ---- data -> entries -----------------------------------------------------------------
def _natkey(s):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s.casefold())]


def _first(name):
    return name.split()[0].casefold()


def _entries(pb):
    rows = []
    for r in pb.live("Pens & Markers"):
        name = NAME_OVERRIDE.get(r["id"], re.sub(r"\s*\(model not given\)", "", r["item"]))
        rows.append({"id": r["id"], "name": name, "sku": r["sku"].strip(), "type": _type(r)})
    counts = {}
    for x in rows:
        counts[_first(x["name"])] = counts.get(_first(x["name"]), 0) + 1
    for x in rows:
        b = _first(x["name"])
        if counts[b] >= GROUP_MIN:
            x["brand"] = b
            rest = x["name"].split(None, 1)[1] if " " in x["name"] else x["name"]
            if rest.split()[0].isalpha() and rest[0].islower() and len(rest.split()[0]) > 1:
                rest = rest[0].upper() + rest[1:]
            x["label"] = rest
            x["key"] = _natkey(b) + [" "] + _natkey(rest)
        else:
            x["brand"] = None
            x["label"] = x["name"]
            x["key"] = _natkey(x["name"])
    rows.sort(key=lambda x: (x["brand"] or x["name"].casefold(), x["key"]))
    brand_disp = {}
    for x in rows:
        if x["brand"] and x["brand"] not in brand_disp:
            brand_disp[x["brand"]] = x["name"].split()[0].upper()
    out, cur_letter, cur_brand, li = [], None, None, -1
    for x in rows:
        letter = x["name"][0].upper()
        if letter != cur_letter:
            li += 1
            cur_letter, cur_brand = letter, None
            out.append({"kind": "letter", "letter": letter, "color": LETTER_COLORS[li % len(LETTER_COLORS)]})
        color = out[-1]["color"] if out[-1]["kind"] == "letter" else _color_of(out, letter)
        if x["brand"] != cur_brand:
            cur_brand = x["brand"]
            if cur_brand:
                out.append({"kind": "brand", "brand": brand_disp[cur_brand], "letter": letter, "color": color})
        out.append({"kind": "item", **x, "letter": letter, "color": color,
                    "brand_disp": brand_disp.get(x["brand"])})
    return out


def _color_of(out, letter):
    for e in reversed(out):
        if e["kind"] == "letter" and e["letter"] == letter:
            return e["color"]
    return LETTER_COLORS[0]


def _columns(entries, ncols=4):
    """Split entries into ncols columns with the smallest row cap that fits.  A column never
    ends on a header; a column that starts mid-letter / mid-brand gets '(continued)' headers."""
    for cap in range(math.ceil(len(entries) / ncols), len(entries) + 1):
        cols, i = [], 0
        ok = True
        while i < len(entries):
            if len(cols) == ncols:
                ok = False
                break
            col = []
            if entries[i]["kind"] != "letter":
                prev = entries[:i]
                lt = next(e for e in reversed(prev) if e["kind"] == "letter")
                col.append({**lt, "cont": True})
                if entries[i]["kind"] == "item" and entries[i]["brand"]:
                    br = next(e for e in reversed(prev) if e["kind"] == "brand")
                    col.append({**br, "cont": True})
            while i < len(entries) and len(col) < cap:
                col.append(entries[i])
                i += 1
            while col and col[-1]["kind"] != "item":       # don't strand a header
                col.pop()
                i -= 1
            if not any(e["kind"] == "item" for e in col):
                ok = False
                break
            cols.append(col)
        if ok:
            while len(cols) < ncols:
                cols.append([])
            return cols, cap
    raise L.SheetDataError("pens sheet: could not lay out columns")


# ---- docx helpers ---------------------------------------------------------------------
def _clear_borders(*cells):
    for c in cells:
        L.set_cell_borders(c, top=None, bottom=None, left=None, right=None)


def _cell_lr(cell, left_in, right_in):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for side, v in (("left", left_in), ("right", right_in)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(v * 1440)))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tcPr.append(m)


def _fit_sizes(label, sku, avail_in):
    name_pt, sku_pt = NAME_PT, SKU_PT
    while True:
        w = _text_w_in(label, name_pt) + (_text_w_in("  " + sku, sku_pt) if sku else 0)
        if w * 1.04 <= avail_in:
            return name_pt, sku_pt
        if name_pt > NAME_MIN_PT:
            name_pt -= 0.5
        elif sku_pt > SKU_MIN_PT:
            sku_pt -= 0.5
        else:
            return name_pt, sku_pt


def _put_letter(row, ci, e):
    c = row.cells[ci].merge(row.cells[ci + 1])
    L.shade(c, e["color"])
    _clear_borders(c)
    p = L.cell_text(c, "", align="left")
    p.paragraph_format.left_indent = Pt(2)
    L.add_run(p, e["letter"], font=L.SANS, size=LETTER_PT, bold=True, color=WHITE)
    if e.get("cont"):
        L.add_run(p, "   continued", font=L.SANS, size=9, italic=True, color=WHITE)


def _put_brand(row, ci, e):
    c = row.cells[ci].merge(row.cells[ci + 1])
    _clear_borders(c)
    L.set_cell_borders(c, bottom=(10, e["color"]))
    p = L.cell_text(c, "", valign="bottom")
    L.add_run(p, e["brand"], font=L.SANS, size=BRAND_PT, bold=True,
              color=RGBColor.from_string(e["color"]), char_spacing_pt=1.2)
    if e.get("cont"):
        L.add_run(p, "  (cont.)", font=L.SANS, size=8.5, color=SKU_GRAY)


def _put_item(row, ci, e, pb):
    nc, pc = row.cells[ci], row.cells[ci + 1]
    indent = INDENT_IN if e["brand"] else 0.0
    name_pt, sku_pt = _fit_sizes(e["label"], e["sku"], NAME_W - 2 * CELL_LR - indent - 0.10)
    p = L.cell_text(nc, "")
    pf = p.paragraph_format
    from docx.shared import Inches
    pf.left_indent = Inches(indent)
    L.add_run(p, e["label"], font=L.SANS, size=name_pt)
    if e["sku"]:
        L.add_run(p, "  " + e["sku"], font=L.SANS, size=sku_pt, color=SKU_GRAY)
    L.cell_text(pc, pb.money(e["id"], dollar=False, cents_style="leading"), align="right",
                font=L.SANS, size=PRICE_PT, bold=True)
    _clear_borders(nc, pc)
    for c in (nc, pc):
        L.set_cell_borders(c, bottom=RULE)
    if e["type"]:
        for c in (nc, pc):
            L.shade(c, TINTS[e["type"]][0])


def _spacer(doc, pt=4, page_break=False):
    p = L.para(doc, line=pt)
    if page_break:
        p.add_run().add_break(WD_BREAK.PAGE)
    r = p.add_run("")
    L.style_run(r, size=1)
    return p


def _header(doc, side, nsides, first_word, last_word, other, letters_here, all_letters, colors):
    t = L.table(doc, 2, 2, widths_in=[4.45, 3.25], borders="none")
    L.set_row_height(t.rows[0], 0.46, "exact")
    L.set_row_height(t.rows[1], 0.30, "exact")
    a, b = t.rows[0].cells
    L.cell_text(a, "PENS & MARKERS", font=L.SANS, size=22, bold=True, valign="center")
    p = L.cell_text(b, "", align="right", valign="center")
    L.add_run(p, f"{first_word} → {last_word}", font=L.SANS, size=14, bold=True)
    L.add_run(p, f"\nSide {side} of {nsides} · flip for {other}", font=L.SANS, size=8.5, color=SKU_GRAY)
    c, d = t.rows[1].cells
    p = L.cell_text(c, "", valign="center")
    for i, lt in enumerate(all_letters):
        if i:
            L.add_run(p, " ", font=L.SANS, size=12)
        if lt in letters_here:
            L.add_run(p, f" {lt} ", font=L.SANS, size=12, bold=True, color=WHITE,
                      highlight_hex=colors[lt])
        else:
            L.add_run(p, f" {lt} ", font=L.SANS, size=12, bold=True, color=STRIP_OFF)
    p = L.cell_text(d, "", align="right", valign="center")
    for k in ("hl", "pencil", "erase"):
        fill, label = TINTS[k]
        L.add_run(p, "  ", font=L.SANS, size=8.5)
        L.add_run(p, f" {label} ", font=L.SANS, size=8.5, highlight_hex=fill)
    # thick rule under the header
    for cell in t.rows[1].cells:
        L.set_cell_borders(cell, bottom=(12, "000000"))
    return t


def _body(doc, left, right, row_in, pb):
    n = max(len(left), len(right))
    t = L.table(doc, n, 5, widths_in=[NAME_W, PRICE_W, GAP_W, NAME_W, PRICE_W],
                align="center", borders="none")
    tblPr = t._tbl.tblPr
    old = tblPr.find(qn("w:tblCellMar"))
    if old is not None:
        tblPr.remove(old)
    L.set_cell_margins_table(t, 0.0, CELL_LR)
    for i in range(n):
        row = t.rows[i]
        L.set_row_height(row, row_in, "exact")
        _clear_borders(row.cells[2])
        for ci, half in ((0, left), (3, right)):
            if i >= len(half):
                _clear_borders(row.cells[ci], row.cells[ci + 1])
                continue
            e = half[i]
            if e["kind"] == "letter":
                _put_letter(row, ci, e)
            elif e["kind"] == "brand":
                _put_brand(row, ci, e)
            else:
                _put_item(row, ci, e, pb)
    L.no_split_rows(t)
    return t


def _guide(col_items):
    items = [e for col in col_items for e in col if e["kind"] == "item" and not e.get("cont")]
    w = lambda e: (e["brand_disp"].title() if e["brand"] else e["name"].split()[0].rstrip(","))  # noqa: E731
    return w(items[0]), w(items[-1])


def build(pb: L.Prices):
    entries = _entries(pb)
    cols, cap = _columns(entries, 4)
    usable = 11 - MARGINS[0] - MARGINS[2] - HEADER_IN - SLACK_IN
    row_in = min(ROW_MAX_IN, usable / max(len(c) for c in cols))
    if row_in < ROW_MIN_IN:
        raise L.SheetDataError(f"pens sheet: {cap} rows per column won't fit two pages at a readable size")

    doc = L.new_document(margins_in=MARGINS, base_font=L.SANS)
    colors = {e["letter"]: e["color"] for e in entries if e["kind"] == "letter"}
    all_letters = list(colors)
    pages = [(cols[0], cols[1]), (cols[2], cols[3])]
    guides = [_guide(pg) for pg in pages]
    for side, (left, right) in enumerate(pages, 1):
        if side > 1:
            _spacer(doc, 1, page_break=True)
        here = {e["letter"] for e in left + right}
        other = guides[1] if side == 1 else guides[0]
        _header(doc, side, len(pages), *guides[side - 1], f"{other[0]} → {other[1]}",
                here, all_letters, colors)
        _spacer(doc, 5)
        _body(doc, left, right, row_in, pb)
    L.rev_stamp(doc, pb)
    return doc
